import copy
import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import qbo
from app.main import app
from app.services.persistence.repository import DocumentRepository
from tests.test_export import seed

client = TestClient(app)


class ReadOnlyQBO:
    realm_id = "realm-103"
    def get_vendors(self):
        return [{"id": "30", "name": "Books by Bessie", "currency": "USD"}]
    def get_expense_accounts(self):
        return [{"id": "10", "name": "Subscriptions", "fully_qualified_name": "Expenses:Subscriptions", "currency": "USD"}]
    def get_currency_preferences(self):
        return {"home_currency": "USD", "multicurrency_enabled": False}


@pytest.fixture
def bill(monkeypatch):
    document = seed(status="extracted", qbo_transaction_id=None, qbo_transaction_type=None,
                    qbo_push_state="not_started", qbo_attachment_id=None,
                    extracted_data={"vendor": "OCR Vendor", "invoice_number": "INV-103",
                                    "date": "2026-10-12", "due_date": "2026-10-30",
                                    "total_amount": "110.00", "tax_amount": "10.00", "currency": "USD"})
    monkeypatch.setattr(qbo, "_get_connected_qbo_client", lambda: ReadOnlyQBO())
    return document


def bill_request(document, **extra):
    params = {"document_id": document["id"], "vendor_id": "30", "account_id": "10",
              "expected_revision": document["_revision"], "expected_realm_id": "realm-103", "tax_code": "TAX"}
    params.update(extra)
    return client.get("/api/export/qbo-import", params=params)


def update(document, **values):
    document.update(values)
    DocumentRepository().save(document)


def test_excel_real_workbook_values_and_filters():
    seed()
    response = client.get("/api/export/excel")
    assert response.status_code == 200 and response.headers["content-disposition"].endswith('.xlsx"')
    book = load_workbook(io.BytesIO(response.content))
    sheet = book["Documents"]
    assert sheet["B2"].value == "fatura-İ.pdf" and sheet["G2"].value == "550.00"
    assert sheet["G2"].data_type == "s" and sheet.freeze_panes == "A2"
    assert sheet.max_column == 16 and sheet.max_row == 2
    assert sheet["A1"].font.bold and sheet.auto_filter.ref == "A1:P2"
    book.close()
    filtered = client.get("/api/export/excel?status=failed")
    book = load_workbook(io.BytesIO(filtered.content))
    assert book.active.max_row == 1
    book.close()


@pytest.mark.parametrize("text", ["=1+1", "+SUM(1)", "-cmd", "@SUM(1)"])
def test_excel_never_creates_formula_cells(text):
    seed(filename=text)
    book = load_workbook(io.BytesIO(client.get("/api/export/excel").content))
    assert book.active["B2"].value == text and book.active["B2"].data_type == "s"
    book.close()


@pytest.mark.parametrize("text", ["bad\x00name", "x" * 32768], ids=lambda value: "excel-invalid-value")
def test_excel_unrepresentable_values_are_explicit_errors(text):
    seed(filename=text)
    assert client.get("/api/export/excel").status_code == 422


def test_bill_csv_existing_names_net_tax_currency_and_no_mutation(bill):
    before = copy.deepcopy(DocumentRepository().get(bill["id"]))
    response = bill_request(bill)
    assert response.status_code == 200
    row = next(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert row["Supplier"] == "Books by Bessie" and row["Account"] == "Expenses:Subscriptions"
    assert row["Line Amount"] == "100.00" and row["Line Tax Amount"] == "10.00"
    assert row["Bill no."] == "INV-103" and row["Currency"] == "USD"
    assert row["Due Date"] == "2026-10-30"
    assert response.headers["x-export-qbo-realm"] == "realm-103"
    assert response.headers["cache-control"] == "no-store"
    assert DocumentRepository().get(bill["id"]) == before
    assert bill_request(bill).content == response.content


@pytest.mark.parametrize("field,value", [("invoice_number", None), ("date", None), ("due_date", None),
    ("tax_amount", None), ("tax_amount", "-1"), ("tax_amount", "200"),
    ("total_amount", "-1"), ("due_date", "2026-10-01"), ("invoice_number", "=1+1")])
def test_bill_missing_or_unsafe_fields_block_export(bill, field, value):
    bill["extracted_data"][field] = value
    update(bill)
    assert bill_request(bill).status_code == 422


@pytest.mark.parametrize("values", [{"approved": False}, {"validation": {"is_valid": False}},
    {"qbo_push_state": "uncertain"}, {"qbo_push_state": "in_progress"},
    {"qbo_transaction_id": "already-written"}, {"status": "pushed"}])
def test_bill_review_and_financial_guards(bill, values):
    update(bill, **values)
    assert bill_request(bill).status_code == 409


@pytest.mark.parametrize("params,code", [({"expected_revision": 99}, 409),
    ({"expected_realm_id": "other-realm"}, 409), ({"vendor_id": "unknown"}, 400),
    ({"account_id": "unknown"}, 400), ({"tax_code": ""}, 422)])
def test_bill_revision_realm_ids_and_tax_code(bill, params, code):
    assert bill_request(bill, **params).status_code == code


def test_bill_no_tax_explicit_zero_and_currency_guard(bill):
    bill["extracted_data"]["tax_amount"] = "0.00"
    update(bill)
    assert bill_request(bill, tax_code="").status_code == 200
    bill["extracted_data"]["currency"] = "EUR"
    update(bill)
    assert bill_request(bill).status_code == 400


def test_receipt_is_not_misrepresented_as_unpaid_bill(bill):
    update(bill, document_type="receipt")
    assert bill_request(bill).status_code == 400


def test_bill_revision_changed_during_qbo_reads(bill, monkeypatch):
    class Concurrent(ReadOnlyQBO):
        def get_vendors(self):
            update(bill, approved=False)
            return super().get_vendors()
    monkeypatch.setattr(qbo, "_get_connected_qbo_client", lambda: Concurrent())
    assert bill_request(bill).status_code == 409


@pytest.mark.auth_live
def test_new_formats_require_login_and_own_company(monkeypatch):
    from app.services.persistence.auth_repository import AuthRepository
    from manage_users import initialize_company
    password = "Formats-password-103"
    own = initialize_company("default", "Own")
    initialize_company("other", "Other")
    AuthRepository().create_user(own, "formats@example.test", password, "User", "user")
    seed()
    seed("doc_other", company="other", filename="other-secret.pdf")
    visitor = TestClient(app)
    assert visitor.get("/api/export/excel").status_code == 401
    assert visitor.get("/api/export/qbo-import").status_code == 401
    assert visitor.post("/api/auth/login", json={"email": "formats@example.test", "password": password}).status_code == 200
    book = load_workbook(io.BytesIO(visitor.get("/api/export/excel").content))
    assert book.active.max_row == 2 and book.active["A2"].value == "doc_export"
    book.close()
    assert visitor.get("/api/export/excel?company_id=other").status_code == 403
    params = {"document_id": "doc_other", "vendor_id": "30", "account_id": "10",
              "expected_revision": 1, "expected_realm_id": "realm-103"}
    assert visitor.get("/api/export/qbo-import", params=params).status_code == 404
