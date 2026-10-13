"""Read-only, company-scoped reports from persisted workflow fields."""
import csv
import io
import json
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy import select
from sqlalchemy.exc import SQLAlchemyError

from app.core import database
from app.core.auth import Principal, company_key, get_current_user
from app.models.database import Document
from app.schemas.documents import (
    BankStatementData, BankTransaction, CheckData, CheckImage, DocumentType,
    InvoiceValidationResult, LineItem, ProcessingStatus, ReceiptData,
)

router = APIRouter()
MAX_DOCUMENTS = 1000
CSV_COLUMNS = [
    "ID", "Filename", "Upload Date", "Status", "Document Type", "Vendor",
    "Amount", "Currency", "Date", "Category", "Approved", "QBO ID",
    "QBO Type", "QBO Status", "Attachment ID", "Attachment Status",
]


def _project(data, model, nested=None):
    """Keep actual extraction fields, without server paths or raw OCR text."""
    if not isinstance(data, dict):
        return None
    result = {key: value for key, value in data.items()
              if key in model.model_fields and key not in {"image_path", "raw_text"}}
    for key, (child_model, many, children) in (nested or {}).items():
        if key not in result:
            continue
        value = result[key]
        if many:
            result[key] = [_project(item, child_model, children)
                           for item in value if isinstance(item, dict)] if isinstance(value, list) else []
        else:
            result[key] = _project(value, child_model, children)
    return result


def _extracted(data, kind):
    if kind in {"receipt", "invoice", "bill"}:
        return _project(data, ReceiptData, {"line_items": (LineItem, True, None)})
    if kind == "bank_statement":
        return _project(data, BankStatementData, {
            "transactions": (BankTransaction, True, {"check_image": (CheckData, False, None)}),
            "check_images": (CheckImage, True, None),
        })
    if kind == "check":
        return _project(data, CheckData)
    return None


def _timestamp(value):
    if value is None:
        return None
    return value.replace(tzinfo=timezone.utc).isoformat() if value.tzinfo is None else value.isoformat()


def report(
    company_id: str | None = Query(None, min_length=1, max_length=100),
    status: ProcessingStatus | None = None,
    document_type: DocumentType | None = None,
    days: int = Query(30, ge=1, le=365),
    user: Principal = Depends(get_current_user),
):
    scope = company_key(company_id)
    now = datetime.now(timezone.utc)
    filters = [Document.company_id == user.company_id, Document.api_id.is_not(None),
               Document.created_at >= (now - timedelta(days=days)).replace(tzinfo=None)]
    if status is not None:
        filters.append(Document.status == status.value)
    if document_type is not None:
        filters.append(Document.document_type == document_type.value)
    # Project individual JSON fields: source bytes, cropped images, tokens and
    # internal push contexts never enter the result or the Python query rows.
    fields = ["extracted_data", "validation", "approved", "qbo_transaction_id",
              "qbo_transaction_type", "qbo_push_state", "qbo_attachment_id",
              "qbo_attachment_push_status"]
    query = select(Document.api_id, Document.filename, Document.content_type,
                   Document.created_at, Document.processed_at, Document.status,
                   Document.document_type,
                   *((Document.workflow_state[field].as_boolean() if field == "approved"
                      else Document.workflow_state[field]).label(field) for field in fields))
    try:
        with database.SessionLocal() as session:
            rows = session.execute(query.where(*filters).order_by(
                Document.created_at.desc(), Document.api_id.desc()).limit(MAX_DOCUMENTS + 1)).mappings().all()
    except SQLAlchemyError:
        raise HTTPException(503, "Export storage is unavailable") from None
    if len(rows) > MAX_DOCUMENTS:
        raise HTTPException(413, "Export exceeds 1000 documents; narrow days, status or document_type")
    items = []
    for row in rows:
        items.append({
            "id": row["api_id"], "filename": row["filename"], "content_type": row["content_type"],
            "created_at": _timestamp(row["created_at"]), "processed_at": _timestamp(row["processed_at"]),
            "status": row["status"], "document_type": row["document_type"],
            "approved": row["approved"] is True,
            "extracted_data": _extracted(row["extracted_data"], row["document_type"]),
            "validation": _project(row["validation"], InvoiceValidationResult),
            "qbo_transaction_id": row["qbo_transaction_id"],
            "qbo_transaction_type": row["qbo_transaction_type"],
            "qbo_push_state": row["qbo_push_state"],
            "qbo_attachment_id": row["qbo_attachment_id"],
            "qbo_attachment_push_status": row["qbo_attachment_push_status"],
        })
    return {"exported_at": now.isoformat(), "total_documents": len(items),
            "filters": {"company_id": scope, "status": status, "document_type": document_type, "days": days},
            "documents": items}


def _csv_cell(value):
    if value is None:
        return ""
    text = str(value)
    # Quoting CSV is insufficient to prevent spreadsheet formula execution.
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")):
        return "'" + text
    return text


def _headers(extension):
    filename = f"receipt-ai-export-{datetime.now(timezone.utc):%Y%m%d}.{extension}"
    return {"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"}


def _summary_values(document):
    extracted = document["extracted_data"] or {}
    return [document["id"], document["filename"], document["created_at"], document["status"],
            document["document_type"], extracted.get("vendor"), extracted.get("total_amount"),
            extracted.get("currency"), extracted.get("date"), extracted.get("category_suggestion"),
            document["approved"], document["qbo_transaction_id"], document["qbo_transaction_type"],
            document["qbo_push_state"], document["qbo_attachment_id"], document["qbo_attachment_push_status"]]


@router.get("/json")
def export_json(data: dict = Depends(report)):
    return Response(json.dumps(data, ensure_ascii=False, indent=2), media_type="application/json",
                    headers=_headers("json"))


@router.get("/csv")
def export_csv(data: dict = Depends(report)):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(CSV_COLUMNS)
    for document in data["documents"]:
        writer.writerow([_csv_cell(value) for value in _summary_values(document)])
    # UTF-8 BOM preserves Turkish characters in spreadsheet applications.
    return Response(("\ufeff" + output.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                    headers=_headers("csv"))


@router.get("/excel")
def export_excel(data: dict = Depends(report)):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Documents"
    for row_number, values in enumerate([CSV_COLUMNS] + [_summary_values(d) for d in data["documents"]], 1):
        for column, value in enumerate(values, 1):
            text = "" if value is None else str(value)
            if len(text) > 32767:
                raise HTTPException(422, "An export value exceeds the Excel cell text limit")
            if ILLEGAL_CHARACTERS_RE.search(text):
                raise HTTPException(422, "An export value contains characters unsupported by Excel; use JSON")
            cell = sheet.cell(row_number, column, text)
            # Keep precision and IDs exactly; never interpret extracted text as a formula.
            cell.data_type = "s"
            if row_number == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="245BE5")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(50, max(12, max(len(str(c.value or "")) for c in column) + 2))
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=_headers("xlsx"))


@router.get("/qbo-import")
def export_qbo_import(
    document_id: str = Query(..., min_length=1, max_length=100),
    vendor_id: str = Query(..., min_length=1, max_length=100),
    account_id: str = Query(..., min_length=1, max_length=100),
    expected_revision: int = Query(..., ge=1),
    expected_realm_id: str = Query(..., min_length=1, max_length=100),
    tax_code: str = Query("", max_length=100),
):
    """One approved invoice -> Bill import draft; never records a QBO write."""
    from app.api.documents import _load_document, _ensure_workflow_editable
    from app.api.qbo import _get_connected_qbo_client
    from app.services.qbo.currency_guard import check_document_currency
    from app.services.qbo.push_guard import ensure_invoice_can_start_push, QBOPushGuardError
    document = _load_document(document_id)
    if document.get("_revision") != expected_revision:
        raise HTTPException(409, "Document changed; reload before exporting")
    if document.get("document_type") != "invoice":
        raise HTTPException(400, "Bill CSV draft supports invoices only; paid receipts are Expenses")
    _ensure_workflow_editable(document)
    if document.get("status") != "extracted":
        raise HTTPException(409, "Bill CSV requires an extracted, approved invoice")
    try:
        ensure_invoice_can_start_push(document)
    except QBOPushGuardError as exc:
        raise HTTPException(409, str(exc)) from None
    from pydantic import ValidationError
    try:
        receipt = ReceiptData.model_validate(document["extracted_data"])
    except ValidationError:
        raise HTTPException(422, "Invoice extraction fields are invalid") from None
    if not receipt.invoice_number or not receipt.invoice_number.strip() or not receipt.date or not receipt.due_date:
        raise HTTPException(422, "Bill CSV requires invoice_number, date and due_date")
    amount, tax = receipt.total_amount, receipt.tax_amount
    if amount is None or tax is None or not amount.is_finite() or not tax.is_finite() or amount < 0 or tax < 0 or tax > amount:
        raise HTTPException(422, "Bill CSV requires explicit nonnegative total_amount and tax_amount (0 for no tax)")
    if receipt.due_date < receipt.date:
        raise HTTPException(422, "Due date cannot precede bill date")
    if tax > 0 and not tax_code.strip():
        raise HTTPException(422, "Taxed bills require an explicit tax code for QBO import mapping")
    client = _get_connected_qbo_client()
    if client.realm_id != expected_realm_id:
        raise HTTPException(409, "QBO company changed; reload options")
    try:
        vendor = next((v for v in client.get_vendors() if str(v["id"]) == vendor_id and v.get("active") is not False), None)
        account = next((a for a in client.get_expense_accounts() if str(a["id"]) == account_id and a.get("active") is not False), None)
    except Exception:
        raise HTTPException(502, "Could not read QBO vendor/account options") from None
    if not vendor or not account:
        raise HTTPException(400, "Selected vendor or expense account is unavailable")
    check_document_currency(client, receipt, vendor, account)
    supplier = vendor.get("name")
    account_name = account.get("fully_qualified_name") or account.get("name")
    if not supplier or not account_name:
        raise HTTPException(502, "QBO vendor/account names are missing")
    # Unlike report CSVs, do not prefix import values: that changes accounting
    # entity names. Reject formula/control prefixes instead of silently editing.
    values = [receipt.invoice_number.strip(), supplier, receipt.date.isoformat(), receipt.due_date.isoformat(),
              account_name, format(amount - tax, "f"), tax_code.strip(), format(tax, "f"),
              receipt.currency, f"Document {document_id}"]
    if any(_csv_cell(value) != str(value) or any(c in str(value) for c in "\r\n\t\x00") for value in values):
        raise HTTPException(422, "Bill CSV contains unsafe import text; correct the source fields")
    if _load_document(document_id).get("_revision") != expected_revision:
        raise HTTPException(409, "Document changed during export; reload")
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["Bill no.", "Supplier", "Bill Date", "Due Date", "Account", "Line Amount",
                     "Line Tax Code", "Line Tax Amount", "Currency", "Line Description"])
    writer.writerow(values)
    return Response(("\ufeff" + output.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="qbo-bill-import-draft.csv"',
                             "Cache-Control": "no-store", "X-Export-QBO-Realm": client.realm_id})
